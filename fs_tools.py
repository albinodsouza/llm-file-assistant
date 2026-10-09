
from pathlib import Path
from datetime import datetime


def extract_text(filepath):
    """Extract text from PDF, TXT, or DOCX files."""
    path = Path(filepath)
    extension = path.suffix.lower()

    if extension == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    elif extension == ".pdf":
        import fitz

        with fitz.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    elif extension == ".docx":
        from docx import Document

        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs)

    else:
        raise ValueError("Only PDF, TXT, and DOCX files are supported.")


def read_file(filepath: str) -> dict:
    """Read a supported file and return content and metadata."""
    try:
        path = Path(filepath)

        if not path.exists():
            return {"success": False, "error": "File not found"}

        if not path.is_file():
            return {"success": False, "error": "Path is not a file"}

        content = extract_text(path)
        stat = path.stat()

        return {
            "success": True,
            "filename": path.name,
            "filepath": str(path.resolve()),
            "extension": path.suffix.lower(),
            "size_bytes": stat.st_size,
            "modified_date": datetime.fromtimestamp(
                stat.st_mtime
            ).isoformat(timespec="seconds"),
            "content": content
        }

    except Exception as error:
        return {
            "success": False,
            "error": str(error)
        }


def list_files(directory: str, extension: str = None) -> list:
    """List files and metadata, optionally filtering by extension."""
    try:
        folder = Path(directory)

        if not folder.exists() or not folder.is_dir():
            return [{
                "success": False,
                "error": "Directory not found or invalid"
            }]

        if extension:
            extension = extension.lower()
            if not extension.startswith("."):
                extension = "." + extension

        results = []

        for path in sorted(folder.iterdir()):
            if not path.is_file():
                continue

            if extension and path.suffix.lower() != extension:
                continue

            stat = path.stat()

            results.append({
                "name": path.name,
                "filepath": str(path.resolve()),
                "extension": path.suffix.lower(),
                "size_bytes": stat.st_size,
                "modified_date": datetime.fromtimestamp(
                    stat.st_mtime
                ).isoformat(timespec="seconds")
            })

        return results

    except Exception as error:
        return [{"success": False, "error": str(error)}]


def write_file(filepath: str, content: str) -> dict:
    """Write content to a file and create missing directories."""
    try:
        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)

        path.write_text(content, encoding="utf-8")

        return {
            "success": True,
            "filepath": str(path.resolve()),
            "filename": path.name,
            "message": "File written successfully",
            "size_bytes": path.stat().st_size
        }

    except Exception as error:
        return {
            "success": False,
            "error": str(error)
        }


def search_in_file(filepath: str, keyword: str) -> dict:
    """Search case-insensitively and return matches with context."""
    if not keyword or not keyword.strip():
        return {
            "success": False,
            "error": "Keyword cannot be empty"
        }

    result = read_file(filepath)

    if not result["success"]:
        return result

    content = result["content"]
    matches = []
    start = 0

    while True:
        index = content.lower().find(keyword.lower(), start)

        if index == -1:
            break

        left = max(0, index - 60)
        right = min(len(content), index + len(keyword) + 60)

        matches.append({
            "position": index,
            "match": content[index:index + len(keyword)],
            "context": content[left:right].replace("\n", " ")
        })

        start = index + max(1, len(keyword))

    return {
        "success": True,
        "filename": result["filename"],
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches
    }
